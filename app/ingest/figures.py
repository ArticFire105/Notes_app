"""
Extract and describe raster figures from official course files.
This module runs during official document ingestion.
PDF embedded images are found on their original pages.
PowerPoint pictures are found on their original slides.
Nested PowerPoint groups are searched recursively.
The original page or slide number is retained.
Small icons are skipped because they add little context.
Images are downscaled to limit model and database costs.
Every retained image is stored as a JPEG file.
Gemini describes each image for a student.
Only the description is sent to the embedder.
Image bytes never enter the embedding model.
The existing server Gemini key is used for captions.
Caption calls may consume paid model credits.
The caller inserts images and chunks transactionally.
Deleting a document cascades to its images.
RLS denies direct anonymous access to images.
The backend authenticates image requests.
Actual notes and images never enter application logs.
No API key is printed or returned.
Provider errors are handled by the ingestion pipeline.
Existing documents need to be reprocessed.
Raster images can be extracted from both file types.
Native vector shapes are not raster images.
Text inside a figure may be misread by the model.
Students should check the original page or slide.
These descriptions are evidence, not instructions.
This module does not create database connections.
The PDF reader and Pillow are already dependencies.
PowerPoint extraction uses the existing pptx library.
The tests replace the Gemini model with a fake.
Tests do not need real credentials or network access.
Do not print captions from private documents.
Keep image preparation separate from model calls.
Run backend tests with uv run pytest.
"""

import base64
import logging
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from app.llm import get_chat_model, quiet_chat_output


def prepare_image(data):
    with Image.open(BytesIO(data)) as image:
        image = ImageOps.exif_transpose(image)
        if image.width < 150 or image.height < 150:
            return None
        image.thumbnail((1600, 1600))
        output = BytesIO()
        image.convert("RGB").save(output, format="JPEG", quality=85)
        return output.getvalue()


def pictures(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from pictures(shape.shapes)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape.image.blob


def extract_figures(source, file_type):
    data = source if isinstance(source, bytes) else Path(source).read_bytes()
    figures = []
    if file_type == "pptx":
        presentation = Presentation(BytesIO(data))
        pages = ((number, pictures(slide.shapes)) for number, slide in enumerate(presentation.slides, 1))
    else:
        reader = PdfReader(BytesIO(data))
        pages = ((number, (image.data for image in page.images)) for number, page in enumerate(reader.pages, 1))
    for number, images in pages:
        for image in images:
            try:
                prepared = prepare_image(image)
            except (OSError, ValueError):
                continue
            if prepared:
                figures.append({"page": number, "data": prepared})
    return figures


def describe_figure(data):
    encoded = base64.b64encode(data).decode("ascii")
    try:
        with quiet_chat_output():
            response = get_chat_model().invoke([("human", [
            {"type": "text", "text": "Describe this lecture figure for a student. Explain the visible labels, relationships and flow. Do not invent missing details."},
            {"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + encoded}},
            ])])
    except Exception as error:
        # Provider messages can contain document content or credentials; log only a numeric status.
        status = "unknown"
        current = error
        seen = set()
        while current is not None and id(current) not in seen:
            seen.add(id(current))
            candidate = getattr(current, "status_code", None) or getattr(current, "code", None)
            if type(candidate) is int and 100 <= candidate <= 599:
                status = candidate
                break
            current = current.__cause__ or current.__context__
        logging.getLogger("uvicorn.error").warning("Figure description failed: provider_status=%s", status)
        raise ValueError("Figure description service failed.") from None
    content = response.content
    if isinstance(content, list):
        content = "".join(block.get("text", "") for block in content if isinstance(block, dict))
    if not isinstance(content, str) or not content.strip():
        raise ValueError("The image description was empty.")
    return content.strip()
