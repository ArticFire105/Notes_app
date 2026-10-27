"""Exercise PDF loading, atomic ingestion, and CLI errors without live services.
Mock model calls and database contexts while checking stored metadata and rollback.
Use temporary dummy PDFs so no real notes, credentials, or .env are accessed.
"""

from types import SimpleNamespace
from unittest.mock import MagicMock
from uuid import UUID

from pydantic import SecretStr
from pypdf import PdfWriter
import pytest

from app import db
from app.ingest import pdf_loader, pipeline
from scripts import ingest_pdf as cli
from io import BytesIO
from zipfile import BadZipFile
from pptx import Presentation
from pptx.util import Inches
from app.ingest import pptx_loader
from app.ingest.chunker import chunk_pages


def test_loader_preserves_original_pages_and_counts_short_text(tmp_path, monkeypatch):
    source = tmp_path / "dummy.pdf"
    source.write_bytes(b"dummy")
    texts = ["A useful explanation of virtual memory and paging.", None, "Page 3", " " * 40,
             "A useful explanation of address translation and page tables."]
    reader = SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda text=text: text) for text in texts])
    monkeypatch.setattr(pdf_loader, "PdfReader", lambda stream: reader)
    pages, skipped = pdf_loader.load_pdf_pages(source)
    assert [page["page"] for page in pages] == [1, 5]
    assert skipped == 3


def test_real_blank_pdf_counts_every_page_as_skipped(tmp_path):
    source = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_blank_page(width=100, height=100)
    with source.open("wb") as output:
        writer.write(output)
    assert pdf_loader.load_pdf_pages(source) == ([], 2)


@pytest.fixture
def ingestion(monkeypatch):
    monkeypatch.setattr(pipeline, "extract_figures", lambda *args: [])
    monkeypatch.setattr(pipeline, "Settings", lambda: SimpleNamespace(embedding_model="test-model"))
    monkeypatch.setattr(pipeline, "load_pdf_pages", lambda path: ([{"page": 2, "text": "notes"}], 1))
    chunks = [{"page": 2, "chunk_index": i, "content": f"chunk {i}"} for i in range(65)]
    monkeypatch.setattr(pipeline, "chunk_pages", lambda pages: chunks)
    embed = MagicMock(side_effect=lambda texts: [[0.0] * 768 for text in texts])
    monkeypatch.setattr(pipeline, "embed_documents", embed)
    connection_context = MagicMock()
    connection = connection_context.return_value.__enter__.return_value
    cursor = connection.cursor.return_value.__enter__.return_value
    document_id = UUID("00000000-0000-0000-0000-000000000001")
    cursor.fetchone.side_effect = [(1, 2), None, (document_id,)]
    monkeypatch.setattr(pipeline, "get_conn", connection_context)
    return SimpleNamespace(
        embed=embed, context=connection_context, cursor=cursor,
        document_id=document_id, chunks=chunks,
    )


def test_ingestion_batches_and_stores_scope_and_original_page_count(ingestion, capsys):
    result = pipeline.ingest_pdf("dummy.pdf", "os", 2, "Memory")
    assert result == ingestion.document_id
    assert [len(call.args[0]) for call in ingestion.embed.call_args_list] == [32, 32, 1]
    calls = ingestion.cursor.execute.call_args_list
    assert calls[0].args[1] == ("OS", 2)
    assert "for update of modules" in calls[0].args[0]
    assert calls[2].args[1] == (1, 2, "official", "Memory", 2)
    rows = ingestion.cursor.executemany.call_args.args[1]
    assert len(rows) == 65
    assert rows[0][:7] == (result, 1, 2, "official", 2, 0, "chunk 0")
    assert rows[0][7].dimensions() == 768
    assert rows[0][8] == "test-model"
    ingestion.context.return_value.__exit__.assert_called_once_with(None, None, None)
    summary = capsys.readouterr().out
    assert "Pages read: 2" in summary and "chunks stored: 65" in summary
    assert "pages skipped (no usable text): 1" in summary and "seconds taken:" in summary


def test_duplicate_is_refused_without_writes(ingestion):
    ingestion.cursor.fetchone.side_effect = [(1, 2), (ingestion.document_id,)]
    with pytest.raises(pipeline.IngestError, match="--replace"):
        pipeline.ingest_pdf("dummy.pdf", "OS", 2, "Memory")
    ingestion.cursor.executemany.assert_not_called()
    assert len(ingestion.cursor.execute.call_args_list) == 2


def test_replace_deletes_old_document_in_same_transaction(ingestion):
    ingestion.cursor.fetchone.side_effect = [(1, 2), (ingestion.document_id,), (ingestion.document_id,)]
    pipeline.ingest_pdf("dummy.pdf", "OS", 2, "Memory", replace=True)
    calls = ingestion.cursor.execute.call_args_list
    assert calls[2].args == (
        "delete from documents where module_id = %s and title = %s", (2, "Memory")
    )
    assert "insert into documents" in calls[3].args[0]
    ingestion.context.assert_called_once()


def test_chunk_insert_failure_exits_transaction_with_error(ingestion, capsys):
    ingestion.cursor.fetchone.side_effect = [(1, 2), (ingestion.document_id,), (ingestion.document_id,)]
    ingestion.cursor.executemany.side_effect = RuntimeError("dummy failure")
    with pytest.raises(pipeline.IngestError, match="save document text"):
        pipeline.ingest_pdf("dummy.pdf", "OS", 2, "Memory", replace=True)
    assert ingestion.context.return_value.__exit__.call_args.args[0] is RuntimeError
    assert "chunks stored:" not in capsys.readouterr().out


def test_third_embedding_batch_failure_opens_no_database_transaction(ingestion):
    ingestion.embed.side_effect = [[[0.0] * 768] * 32, [[0.0] * 768] * 32, RuntimeError("dummy failure")]
    with pytest.raises(pipeline.IngestError, match="Text embedding failed.*RuntimeError"):
        pipeline.ingest_pdf("dummy.pdf", "OS", 2, "Memory")
    ingestion.context.assert_not_called()


def test_all_scanned_pages_are_reported_without_model_or_database(ingestion, monkeypatch):
    monkeypatch.setattr(pipeline, "load_pdf_pages", lambda path: ([], 3))
    monkeypatch.setattr(pipeline, "chunk_pages", lambda pages: [])
    with pytest.raises(pipeline.IngestError, match="pages read=3, pages skipped=3"):
        pipeline.ingest_pdf("dummy.pdf", "OS", 2, "Memory")
    ingestion.embed.assert_not_called()
    ingestion.context.assert_not_called()


def test_unknown_scope_is_refused_without_writes(ingestion):
    ingestion.cursor.fetchone.side_effect = [None]
    with pytest.raises(pipeline.IngestError, match="Subject/module not found"):
        pipeline.ingest_pdf("dummy.pdf", "OS", 99, "Memory")
    ingestion.cursor.executemany.assert_not_called()


def test_connection_registers_vectors_and_forwards_error_to_driver(monkeypatch):
    settings = SimpleNamespace(database_url=SecretStr("dummy-db-url"))
    monkeypatch.setattr(db, "Settings", lambda: settings)
    connect = MagicMock()
    register = MagicMock()
    monkeypatch.setattr(db.psycopg, "connect", connect)
    monkeypatch.setattr(db, "register_vector", register)
    with pytest.raises(RuntimeError, match="dummy failure"):
        with db.get_conn() as connection:
            register.assert_called_once_with(connection)
            raise RuntimeError("dummy failure")
    assert connect.return_value.__exit__.call_args.args[0] is RuntimeError


def test_cli_does_not_expose_external_exception(monkeypatch, capsys):
    monkeypatch.setattr("sys.argv", [
        "ingest_pdf.py", "--pdf", "dummy.pdf", "--subject", "OS", "--module", "2", "--title", "Memory"
    ])
    monkeypatch.setattr(cli, "ingest_pdf", MagicMock(side_effect=RuntimeError("dummy-secret")))
    assert cli.main() == 1
    output = capsys.readouterr()
    assert "FAIL ingestion" in output.err
    assert "dummy-secret" not in output.out + output.err


def test_byte_loader_preserves_original_page_numbers(monkeypatch):
    texts = ["A full explanation of paging from the lecture.", "", "A full explanation of page tables."]
    reader = SimpleNamespace(pages=[
        SimpleNamespace(extract_text=lambda text=text: text) for text in texts
    ])
    monkeypatch.setattr(pdf_loader, "PdfReader", lambda stream: reader)
    pages, skipped = pdf_loader.load_pdf_pages(b"dummy PDF bytes")
    assert [page["page"] for page in pages] == [1, 3]
    assert skipped == 1


def test_existing_document_ingestion_reuses_row_and_marks_ready(ingestion):
    ingestion.cursor.fetchone.side_effect = [(1, 2), (ingestion.document_id,)]
    result = pipeline.ingest_pdf(
        b"dummy PDF bytes", "OS", 2, "Memory", document_id=ingestion.document_id
    )
    assert result == ingestion.document_id
    calls = ingestion.cursor.execute.call_args_list
    assert not any("insert into documents" in call.args[0] for call in calls)
    assert calls[1].args[1] == (result, 1, 2, "official")
    assert "status = 'processing'" in calls[1].args[0]
    assert "status = 'ready'" in calls[-2].args[0]
    assert calls[-1].args == ("delete from answer_cache where subject_id = %s", (1,))
    assert calls[-2].args[1] == (2, result)
    assert ingestion.cursor.executemany.call_args.args[1][0][0] == result
    ingestion.context.return_value.__exit__.assert_called_once_with(None, None, None)


def test_existing_document_scope_mismatch_prevents_chunk_writes(ingestion):
    ingestion.cursor.fetchone.side_effect = [(1, 2), None]
    with pytest.raises(pipeline.IngestError, match="Processing document not found"):
        pipeline.ingest_pdf(
            b"dummy PDF bytes", "OS", 2, "Memory", document_id=ingestion.document_id
        )
    ingestion.cursor.executemany.assert_not_called()


def test_existing_document_chunk_failure_rolls_back_without_ready(ingestion):
    ingestion.cursor.fetchone.side_effect = [(1, 2), (ingestion.document_id,)]
    ingestion.cursor.executemany.side_effect = RuntimeError("dummy failure")
    with pytest.raises(pipeline.IngestError, match="save document text"):
        pipeline.ingest_pdf(
            b"dummy PDF bytes", "OS", 2, "Memory", document_id=ingestion.document_id
        )
    assert ingestion.context.return_value.__exit__.call_args.args[0] is RuntimeError
    assert not any(
        "status = 'ready'" in call.args[0]
        for call in ingestion.cursor.execute.call_args_list
    )


def test_cli_keeps_path_based_ingestion(monkeypatch, capsys):
    monkeypatch.setattr("sys.argv", [
        "ingest_pdf.py", "--pdf", "dummy.pdf", "--subject", "OS",
        "--module", "2", "--title", "Memory", "--replace",
    ])
    ingest = MagicMock(return_value="dummy-document-id")
    monkeypatch.setattr(cli, "ingest_pdf", ingest)
    assert cli.main() == 0
    assert str(ingest.call_args.args[0]) == "dummy.pdf"
    assert ingest.call_args.args[1:] == ("OS", 2, "Memory")
    assert ingest.call_args.kwargs == {"replace": True}
    assert "PASS ingestion" in capsys.readouterr().out


def sample_pptx():
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    first.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Paging"
    first.notes_slide.notes_text_frame.text = "The speaker explains how virtual pages map to physical frames."
    presentation.slides.add_slide(presentation.slide_layouts[6])
    third = presentation.slides.add_slide(presentation.slide_layouts[6])
    third.notes_slide.notes_text_frame.text = "These speaker notes are the only text on slide three."
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_pptx_extracts_short_titles_and_speaker_notes_without_renumbering():
    pages, skipped = pptx_loader.load_pptx_pages(sample_pptx())
    assert [page["page"] for page in pages] == [1, 3]
    assert skipped == 1
    assert pages[0]["text"].startswith("Paging\n\nSpeaker notes:\n")
    assert "physical frames" in pages[0]["text"]
    assert "only text on slide three" in pages[1]["text"]


def test_pptx_supports_paths_and_keeps_chunk_slide_numbers(tmp_path):
    source = tmp_path / "lecture.pptx"
    source.write_bytes(sample_pptx())
    pages, skipped = pptx_loader.load_pptx_pages(source)
    chunks = chunk_pages(pages)
    assert [chunk["page"] for chunk in chunks] == [1, 3]
    assert skipped == 1


def test_pptx_extracts_grouped_text_and_tables():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Grouped lecture text"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "Page"
    table.cell(0, 1).text = "Frame"
    table.cell(1, 0).text = "Virtual"
    table.cell(1, 1).text = "Physical"
    output = BytesIO()
    presentation.save(output)
    pages, skipped = pptx_loader.load_pptx_pages(output.getvalue())
    assert skipped == 0
    assert "Grouped lecture text" in pages[0]["text"]
    assert "Page | Frame\nVirtual | Physical" in pages[0]["text"]


def test_pptx_parser_does_not_create_missing_notes_slides(monkeypatch):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    assert not slide.has_notes_slide
    monkeypatch.setattr(pptx_loader, "Presentation", lambda stream: presentation)
    with pytest.raises(pptx_loader.NoReadableSlidesError, match="No readable text found in slides"):
        pptx_loader.load_pptx_pages(b"fake")
    assert not slide.has_notes_slide


def test_pptx_handles_missing_notes_placeholder(monkeypatch):
    slide = SimpleNamespace(shapes=[], has_notes_slide=True,
                            notes_slide=SimpleNamespace(notes_text_frame=None))
    monkeypatch.setattr(pptx_loader, "Presentation", lambda stream: SimpleNamespace(slides=[slide]))
    with pytest.raises(pptx_loader.NoReadableSlidesError, match="No readable text found in slides"):
        pptx_loader.load_pptx_pages(b"fake")


def test_pptx_rejects_corrupt_presentation():
    with pytest.raises(BadZipFile):
        pptx_loader.load_pptx_pages(b"PK this is not a valid PowerPoint archive")


def test_pptx_uses_shared_chunk_embedding_and_storage_pipeline(ingestion, monkeypatch):
    monkeypatch.setattr(pipeline, "chunk_pages", chunk_pages)
    result = pipeline.ingest_document(sample_pptx(), "OS", 2, "Lecture", file_type="pptx")
    assert result == ingestion.document_id
    rows = ingestion.cursor.executemany.call_args.args[1]
    assert [row[4] for row in rows] == [1, 3]
    assert "physical frames" in rows[0][6]
    assert "only text on slide three" in rows[1][6]
    assert all(row[7].dimensions() == 768 for row in rows)
    assert ingestion.embed.call_count == 1
    assert any(call.args[1] == (3, result) for call in ingestion.cursor.execute.call_args_list
               if "status = 'ready'" in call.args[0])


@pytest.mark.parametrize("text,expected", [
    ("  First\x0b  second  ", "First\nsecond"),
    ("A\x00B\x1bC\x7fD\x85E\tF", "A B C D E F"),
    ("A\r\n B\rC\x0cD", "A\nB\nC\nD"),
    (" \x00\x0b\t\x1b\u200b ", ""),
    (" \u03b1   +  \u03b2  \n\n\n  equals \u03b3 ", "\u03b1 + \u03b2\n\nequals \u03b3"),
])
def test_pptx_cleans_controls_and_whitespace(text, expected):
    assert pptx_loader.clean_slide_text(text) == expected


def deck_with_soft_breaks():
    presentation = Presentation()
    for number in range(1, 24):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        if number in {1, 23}:
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "  Virtual memory\x0b  uses   pages  "
        elif number == 7:
            slide.notes_slide.notes_text_frame.text = "  Speaker\x0b notes    about page tables  "
        elif number == 5:
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "  \x0b  \t "
        elif number == 10:
            slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_23_slide_deck_skips_empty_slides_cleans_soft_breaks_and_keeps_citations():
    from app.rag.prompt import build_context
    pages, skipped = pptx_loader.load_pptx_pages(deck_with_soft_breaks())
    assert skipped == 20
    assert [page["page"] for page in pages] == [1, 7, 23]
    assert pages[0]["text"] == "Virtual memory\nuses pages"
    assert pages[1]["text"] == "Speaker notes:\nSpeaker\nnotes about page tables"
    assert all("\x0b" not in page["text"] and "\x00" not in page["text"] for page in pages)
    context = build_context([
        {"subject": "OS", "module": 2, "document": "Lecture", "file_type": "pptx",
         "page": page["page"], "content": page["text"]} for page in pages
    ])
    for number in (1, 7, 23):
        assert f"Slide {number})" in context


@pytest.mark.parametrize("count", [0, 1, 23])
def test_all_empty_pptx_returns_clear_error(count):
    presentation = Presentation()
    for _ in range(count):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "  \x0b \t "
        slide.notes_slide.notes_text_frame.text = " \x0b "
    output = BytesIO()
    presentation.save(output)
    with pytest.raises(pptx_loader.NoReadableSlidesError, match="^No readable text found in slides$"):
        pptx_loader.load_pptx_pages(output.getvalue())


def test_empty_deck_pipeline_stops_before_embeddings_and_database(ingestion, monkeypatch, caplog):
    def no_text(path):
        raise pptx_loader.NoReadableSlidesError(pptx_loader.NO_READABLE_SLIDES)
    monkeypatch.setattr(pipeline, "load_pptx_pages", no_text)
    with pytest.raises(pipeline.IngestError, match="^No readable text found in slides$") as error:
        pipeline.ingest_document(b"fake", "OS", 2, "Lecture", file_type="pptx")
    assert error.value.safe_reason == "No readable text found in slides"
    ingestion.embed.assert_not_called()
    ingestion.context.assert_not_called()
    assert "Traceback (most recent call last)" in caplog.text
    assert "NoReadableSlidesError" in caplog.text


def test_pipeline_logs_real_stack_without_exception_messages_or_secrets(ingestion, caplog):
    ingestion.embed.side_effect = RuntimeError(
        "fake-private-key https://private.invalid?token=fake-token "
        "postgresql://user:fake-password@private/db Authorization: Bearer fake-access"
    )
    with pytest.raises(pipeline.IngestError, match="Text embedding failed.*RuntimeError") as error:
        pipeline.ingest_document(b"fake", "OS", 2, "Lecture")
    assert isinstance(error.value.__cause__, RuntimeError)
    assert "RuntimeError" in caplog.text
    assert "pipeline.py" in caplog.text
    assert "Traceback (most recent call last)" in caplog.text
    for value in ("fake-private-key", "fake-token", "fake-password", "fake-access", "https://", "postgresql://"):
        assert value not in caplog.text
        assert value not in str(error.value)


def test_embedding_failure_logs_original_chunk_index_and_only_80_preview_characters(ingestion, caplog, monkeypatch):
    from app.embeddings import EmbeddingError
    monkeypatch.setattr(pipeline, "load_pptx_pages", lambda path: ([{"page": 2, "text": "notes"}], 1))
    text = "pages map to physical frames " * 5 + "outside preview"
    ingestion.chunks[33]["content"] = text
    ingestion.embed.side_effect = [
        [[0.0] * 768] * 32,
        EmbeddingError("Text embedding incorrectly entered video preprocessing (ImportError).", chunk_index=1),
    ]
    with pytest.raises(pipeline.IngestError, match="Chunk 33:.*video preprocessing.*ImportError") as error:
        pipeline.ingest_document(b"fake", "OS", 2, "Lecture", file_type="pptx")
    assert error.value.safe_reason.startswith("Chunk 33:")
    preview = next(record.getMessage() for record in caplog.records if "preview:" in record.getMessage())
    assert preview.split("preview: ", 1)[1] == text[:80]
    assert "outside preview" not in caplog.text
    ingestion.context.assert_not_called()


def test_pipeline_skips_empty_chunks_before_embedding_without_losing_alignment(ingestion):
    ingestion.chunks[1]["content"] = " \n "
    pipeline.ingest_document(b"fake", "OS", 2, "Lecture")
    rows = ingestion.cursor.executemany.call_args.args[1]
    assert len(rows) == 64
    assert rows[0][5:7] == (0, "chunk 0")
    assert rows[1][5:7] == (2, "chunk 2")
    assert [len(call.args[0]) for call in ingestion.embed.call_args_list] == [32, 32]


def test_chunk_preview_masks_secrets_before_truncating():
    from app.logging_safety import safe_chunk_preview
    assert safe_chunk_preview("pages " * 30 + "token=fake-private-value") == "[redacted: potentially sensitive text]"
    assert "private" not in safe_chunk_preview("See https://private.invalid/path for details")
    assert "PRIVATE" not in safe_chunk_preview("See HTTPS://PRIVATE.INVALID/path for details")


def test_figure_captions_are_embedded_as_text_and_saved_with_image(ingestion, monkeypatch):
    figure = {"page": 2, "data": b"fake-image-bytes"}
    monkeypatch.setattr(pipeline, "extract_figures", lambda *args: [figure])
    monkeypatch.setattr(pipeline, "describe_figure", lambda data: "The diagram shows a page table.")
    pipeline.ingest_pdf("dummy.pdf", "OS", 2, "Memory")
    all_texts = [item for call in ingestion.embed.call_args_list for item in call.args[0]]
    assert all(isinstance(item, str) for item in all_texts)
    assert any("page table" in item for item in all_texts)
    figure_insert = next(call for call in ingestion.cursor.execute.call_args_list if "insert into course_figures" in call.args[0])
    assert figure_insert.args[1][-1] == b"fake-image-bytes"
    assert any("figure_id" in call.args[0] for call in ingestion.cursor.execute.call_args_list)


def test_extract_pptx_picture_and_skip_small_icon():
    from PIL import Image
    from app.ingest.figures import extract_figures
    def image(size):
        output = BytesIO()
        Image.new("RGB", size, "white").save(output, format="PNG")
        output.seek(0)
        return output
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(image((300, 300)), Inches(1), Inches(1))
    slide.shapes.add_picture(image((20, 20)), Inches(2), Inches(2))
    output = BytesIO()
    deck.save(output)
    figures = extract_figures(output.getvalue(), "pptx")
    assert len(figures) == 1
    assert figures[0]["page"] == 1
    assert figures[0]["data"].startswith(b"\xff\xd8")


def test_ocr_cleaning_removes_markup_without_losing_table_text():
    from app.ocr import clean_ocr_text
    result = clean_ocr_text("Heading <table><tr><td>CPU</td><td>Memory</td></tr></table> &lt;div&gt;Text&lt;/div&gt; <script>junk</script>")
    assert "<" not in result and ">" not in result
    assert "CPU" in result and "Memory" in result and "Text" in result
    assert "junk" not in result
