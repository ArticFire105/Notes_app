"""
Handle official PDF/PPTX upload and document status polling.
The upload response is sent before embedding runs.
The background worker records safe failure messages.
Existing CLI-created documents can lack a Storage path.

This module combines Phase 3 uploads with Phase 4 permissions.
The API wraps the existing ingestion and answering code.
Every document route requires a college login.
Only administrators can upload official notes.

Database credentials stay on the backend.
The Storage service role key stays on the backend.
Responses never include configuration objects.
Raw provider errors must not reach clients or logs.

Subjects and modules define the official-notes scope.
Database IDs are distinct from module display numbers.
Uploaded PDFs live in a private notes bucket.
Signed URLs grant temporary access to one stored PDF.

Uploads are checked before background ingestion.
Document status supports polling during ingestion.
Chunks and readiness are committed together.
Background tasks are in-process rather than a durable queue.

Startup reuses the cached Phase 1 embedding model.
Question answers reuse the Phase 2 answer function.
Tests replace model loading and all external calls.
Tests never need real credentials or a real dotenv file.

Upload records use the verified administrator's profile ID.
Run the API from the repository root.
Use uv run pytest for the offline test suite.
"""

from io import BytesIO
from uuid import UUID, uuid4
import shutil
import subprocess
import tempfile
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import Response
from psycopg.rows import dict_row
from pypdf import PdfReader

from app.db import get_conn
import base64
from io import BytesIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.cache import invalidate_scope
from app.auth import require_role
from app.ingest.pipeline import IngestError, ingest_document
from app.ingest.pptx_loader import NO_READABLE_SLIDES, NoReadableSlidesError, load_pptx_pages
from app.logging_safety import log_safe_exception
from app.schemas import DocumentOut
from app.storage import delete_document, download_document, restore_document, signed_url, upload_document

router = APIRouter()
MAX_UPLOAD_BYTES = 40 * 1024 * 1024
PPTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DELETE_FAILED = "Could not complete document deletion. Please retry or check backend logs."
DELETE_RECOVERY_FAILED = "Delete failed and the file could not be restored. Contact the administrator and check backend logs."
DOCUMENT_COLUMNS = """
    d.id, d.title, m.number as module, d.page_count, d.status, d.created_at, d.error,
    case when lower(d.storage_path) like '%%.pptx' then 'pptx' else 'pdf' end as file_type
"""


def get_document(document_id: UUID):
    with get_conn() as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            cursor.execute(
                "select " + DOCUMENT_COLUMNS + """, d.storage_path
                from documents d left join modules m on m.id = d.module_id
                where d.id = %s
                """,
                (document_id,),
            )
            document = cursor.fetchone()
    if document is None:
        raise HTTPException(404, "Document not found.")
    return document


def finish_upload(file_bytes, document_id, subject_code, module_number, title, file_type="pdf"):
    try:
        ingest_document(file_bytes, subject_code, module_number, title,
                        document_id=document_id, file_type=file_type)
    except Exception as error:
        if isinstance(error, IngestError):
            reason = error.safe_reason
        else:
            reason = "Document ingestion failed. Check backend logs for details."
            log_safe_exception(error, reason)
        try:
            with get_conn() as connection:
                with connection.cursor() as cursor:
                    cursor.execute(
                        "update documents set status = 'failed', error = %s where id = %s",
                        (reason, document_id),
                    )
        except Exception as status_error:
            # If the DB itself is unavailable, status cannot be persisted yet.
            log_safe_exception(status_error, "Could not record ingestion failure status.")
            return


@router.get("/modules/{module_id}/documents", response_model=list[DocumentOut])
def list_documents(module_id: int):
    with get_conn() as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            cursor.execute("select id from modules where id = %s", (module_id,))
            if cursor.fetchone() is None:
                raise HTTPException(404, "Module not found.")
            cursor.execute(
                "select " + DOCUMENT_COLUMNS + """
                from documents d left join modules m on m.id = d.module_id
                where d.module_id = %s and d.source_type = 'official'
                order by d.created_at desc
                """,
                (module_id,),
            )
            return cursor.fetchall()


@router.post("/documents", response_model=DocumentOut, status_code=202)
async def create_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(),
    subject_id: int = Form(gt=0),
    module_id: int = Form(gt=0),
    title: str = Form(min_length=1, max_length=500),
    user=Depends(require_role("admin")),
):
    title = title.strip()
    if not title:
        raise HTTPException(400, "Provide a nonempty document title.")
    # Read only enough to reject an oversized file before Storage or embedding.
    try:
        file_bytes = await file.read(MAX_UPLOAD_BYTES + 1)
    finally:
        await file.close()
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "File must be 40 MB or smaller.")
    filename = (file.filename or "").lower()
    if filename.endswith(".ppt"):
        raise HTTPException(400, "Old .ppt files are not supported. Save as .pptx and upload again.")
    if filename.endswith(".pptx"):
        if (file.content_type not in {PPTX_CONTENT_TYPE, "application/octet-stream", "application/zip"}
                or not file_bytes.startswith(b"PK")):
            raise HTTPException(400, "Upload a valid .pptx PowerPoint file.")
        file_type = "pptx"
    elif (filename.endswith(".pdf") and file.content_type == "application/pdf"
          and file_bytes.startswith(b"%PDF-")):
        file_type = "pdf"
    else:
        raise HTTPException(400, "Upload a PDF or .pptx PowerPoint file.")
    return await run_in_threadpool(
        queue_document, background_tasks, file_bytes, subject_id, module_id, title, user["id"], file_type
    )


def queue_document(background_tasks, file_bytes, subject_id, module_id, title, uploaded_by, file_type="pdf"):
    try:
        if file_type == "pptx":
            pages, skipped = load_pptx_pages(file_bytes)
            if not pages and not skipped:
                raise ValueError("Empty presentation")
        else:
            reader = PdfReader(BytesIO(file_bytes))
            if reader.is_encrypted or len(reader.pages) == 0:
                raise ValueError("Unreadable PDF")
    except NoReadableSlidesError as error:
        log_safe_exception(error, NO_READABLE_SLIDES)
        raise HTTPException(400, NO_READABLE_SLIDES) from None
    except Exception as error:
        log_safe_exception(error, "Could not read uploaded document.")
        raise HTTPException(400, "Upload a readable, unencrypted PDF or .pptx presentation with slides.") from None
    with get_conn() as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            cursor.execute("select code from subjects where id = %s", (subject_id,))
            subject = cursor.fetchone()
            if subject is None:
                raise HTTPException(404, "Subject not found.")
            cursor.execute(
                "select number from modules where id = %s and subject_id = %s",
                (module_id, subject_id),
            )
            module = cursor.fetchone()
            if module is None:
                raise HTTPException(404, "Module not found in this subject.")
    document_id = uuid4()
    path = f"official/{subject['code']}/module-{module['number']}/{document_id}.{file_type}"
    upload_document(file_bytes, path)
    try:
        with get_conn() as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                cursor.execute(
                    """
                    insert into documents
                        (id, subject_id, module_id, source_type, title, storage_path, status, uploaded_by)
                    values (%s, %s, %s, 'official', %s, %s, 'processing', %s)
                    returning id, title, page_count, status, created_at, error
                    """,
                    (document_id, subject_id, module_id, title, path, uploaded_by),
                )
                document = cursor.fetchone()
    except Exception:
        try:
            delete_document(path)
        except Exception:
            pass
        raise HTTPException(500, "Could not record the uploaded document.") from None
    document["module"] = module["number"]
    document["file_type"] = file_type
    background_tasks.add_task(
        finish_upload, file_bytes, document_id, subject["code"], module["number"], title, file_type
    )
    return document


@router.get("/documents/{document_id}", response_model=DocumentOut)
def document_status(document_id: UUID):
    return get_document(document_id)


@router.delete("/documents/{document_id}")
def remove_document(document_id: UUID, user=Depends(require_role("admin"))):
    file_bytes = None
    storage_path = None
    storage_attempted = False
    try:
        with get_conn() as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                cursor.execute("select module_id from documents where id = %s", (document_id,))
                scope = cursor.fetchone()
                if scope is None:
                    raise HTTPException(404, "Document not found.")
                # Match ingestion's module -> document lock order to avoid deadlocks.
                if scope["module_id"] is not None:
                    cursor.execute("select id from modules where id = %s for update", (scope["module_id"],))
                cursor.execute(
                    "select subject_id, storage_path from documents where id = %s for update",
                    (document_id,),
                )
                document = cursor.fetchone()
                if document is None:
                    raise HTTPException(404, "Document not found.")
                storage_path = document["storage_path"]
                # Keep an in-memory copy until commit so a handled failure can restore the file.
                if storage_path:
                    file_bytes = download_document(storage_path)
                cursor.execute("delete from chunks where document_id = %s", (document_id,))
                invalidate_scope(document["subject_id"], cursor)
                cursor.execute("delete from documents where id = %s", (document_id,))
                # All SQL operations must succeed before attempting the external removal.
                if storage_path:
                    storage_attempted = True
                    delete_document(storage_path)
        return {"deleted": True, "id": str(document_id)}
    except HTTPException:
        raise
    except Exception as error:
        log_safe_exception(error, "Document deletion failed.")
        if storage_attempted and file_bytes is not None:
            try:
                restore_document(file_bytes, storage_path)
            except Exception as restore_error:
                log_safe_exception(restore_error, "Could not restore document file after failed deletion.")
                raise HTTPException(503, DELETE_RECOVERY_FAILED) from None
        raise HTTPException(503, DELETE_FAILED) from None


@router.get("/documents/{document_id}/pdf-url")
@router.get("/documents/{document_id}/file-url")
def document_pdf_url(document_id: UUID):
    document = get_document(document_id)
    if not document["storage_path"]:
        raise HTTPException(404, "This document has no stored file; upload it through the API.")
    return {"url": signed_url(document["storage_path"])}



@router.get("/documents/{document_id}/slides/{number}")
def document_slide(document_id: UUID, number: int):
    document = get_document(document_id)
    path = document.get("storage_path")
    if not path or not path.lower().endswith(".pptx"):
        raise HTTPException(404, "Presentation not found.")
    try:
        presentation = Presentation(BytesIO(download_document(path)))
        if number < 1 or number > len(presentation.slides):
            raise HTTPException(404, "Slide not found.")
        width, height = presentation.slide_width, presentation.slide_height
        shapes = []
        for shape in presentation.slides[number - 1].shapes:
            item = {"left": shape.left / width * 100, "top": shape.top / height * 100,
                    "width": shape.width / width * 100, "height": shape.height / height * 100}
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                if image.content_type not in {"image/png", "image/jpeg", "image/gif"}:
                    continue
                item.update(kind="image", data=base64.b64encode(image.blob).decode("ascii"), mime=image.content_type)
            elif shape.has_text_frame and shape.text_frame.text.strip():
                item.update(kind="text", text=shape.text_frame.text)
            else:
                continue
            shapes.append(item)
        return {"number": number, "total": len(presentation.slides),
                "aspect": width / height, "shapes": shapes}
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(503, "Could not display this slide.") from None


def convert_presentation(file_bytes):
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        for candidate in (Path("C:/Program Files/LibreOffice/program/soffice.exe"),
                          Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe")):
            if candidate.is_file():
                executable = str(candidate)
                break
    if not executable:
        raise HTTPException(503, "Install LibreOffice on the backend to preview PowerPoint files.")
    try:
        # A separate profile avoids collisions with the desktop app and other previews.
        with tempfile.TemporaryDirectory(prefix="notes-preview-") as folder:
            root = Path(folder)
            source = root / "presentation.pptx"
            source.write_bytes(file_bytes)
            profile = (root / "profile").as_uri()
            subprocess.run([executable, "--headless", f"-env:UserInstallation={profile}",
                            "--convert-to", "pdf", "--outdir", str(root), str(source)],
                           check=True, timeout=120, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            result = (root / "presentation.pdf").read_bytes()
            if not result.startswith(b"%PDF-"):
                raise ValueError("Invalid conversion")
            return result
    except Exception:
        raise HTTPException(503, "Could not convert this presentation for preview.") from None


@router.get("/documents/{document_id}/preview")
def presentation_preview(document_id: UUID):
    document = get_document(document_id)
    path = document.get("storage_path")
    if not path or not path.lower().endswith(".pptx"):
        raise HTTPException(404, "Presentation not found.")
    try:
        pdf = convert_presentation(download_document(path))
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(503, "Could not load this presentation for preview.") from None
    return Response(pdf, media_type="application/pdf", headers={"Cache-Control": "no-store"})


@router.get("/figures/{figure_id}")
def course_figure(figure_id: UUID):
    with get_conn() as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            cursor.execute("""select f.image_data, f.content_type from course_figures f
                              join documents d on d.id = f.document_id
                              where f.id = %s and d.source_type = 'official' and d.status = 'ready'""", (figure_id,))
            image = cursor.fetchone()
    if image is None:
        raise HTTPException(404, "Figure not found.")
    return Response(bytes(image["image_data"]), media_type=image["content_type"], headers={"Cache-Control": "private, no-store"})
